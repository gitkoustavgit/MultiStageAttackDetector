"""
build_full_report.py

Master script to update and assemble BCSE497J Project I Report.docx.
Integrates:
- report_updates_1_4.py (Updates Abstract, Ch2 literature review with all 30 citations,
                         Ch3 requirements, feasibility, specs, and Ch4 architecture & DFD captions)
- report_chapter_5.py   (Appends Chapter 5: Implementation and Experimental Methodology with Fig 8)
- report_chapter_6.py   (Appends Chapter 6: Results and Evaluation with Figs 9, 10, 11)
- report_chapters_7_8_9.py (Appends Chapters 7, 8, and Chapter 9 References [1]-[30])
- Table of Contents updater (Rebuilds Table 2 with all 9 chapters & subsections)
- Global Black Font Enforcement (Sets RGBColor(0,0,0) across every run in document)
"""

import os
import sys
import shutil
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from report_updates_1_4 import apply_updates_ch1_to_4
from report_chapter_5 import append_chapter_5
from report_chapter_6 import append_chapter_6
from report_chapters_7_8_9 import append_chapters_7_8_9

DOC_PATH = r"C:\Users\Srikar Bharadwaj\Desktop\Projects\Context-Aware AI Detection of Multi-Stage Injection Attacks\BCSE497J Project I Report.docx"
BACKUP_PATH = r"C:\Users\Srikar Bharadwaj\Desktop\Projects\Context-Aware AI Detection of Multi-Stage Injection Attacks\BCSE497J Project I Report.docx.bak"

BLACK = RGBColor(0, 0, 0)
FONT_NAME = "Times New Roman"

def set_run_font(run, size_pt=11, bold=False, italic=False):
    run.font.name = FONT_NAME
    if size_pt:
        run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = BLACK

def format_cell(cell, text="", bold=False, italic=False, size_pt=10.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell.paragraphs[0]
    p.text = ""
    p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    if text:
        r = p.add_run(text)
        set_run_font(r, size_pt=size_pt, bold=bold, italic=italic)

def update_table_of_contents(doc):
    print("Updating Table of Contents (Table 2)...")
    toc = doc.tables[2]
    
    # Remove all existing rows except the header row (row 0)
    for row in list(toc.rows)[1:]:
        toc._tbl.remove(row._tr)
        
    toc_entries = [
        # (Sl.No, Contents, Page No., is_major_chapter)
        ("", "Abstract", "", False),
        ("1.", "INTRODUCTION", "5", True),
        ("", "1.1 Background", "5", False),
        ("", "1.2 Motivation", "6", False),
        ("", "1.3 Scope of the Project", "6", False),
        ("2.", "PROJECT DESCRIPTION AND GOALS", "8", True),
        ("", "2.1 Literature Review", "8", False),
        ("", "2.2 Research Gap", "9", False),
        ("", "2.3 Objectives", "10", False),
        ("", "2.4 Problem Statement", "10", False),
        ("", "2.5 Project Plan", "11", False),
        ("3.", "TECHNICAL SPECIFICATION", "12", True),
        ("", "3.1 Requirements", "12", False),
        ("", "3.1.1 Functional Requirements", "12", False),
        ("", "3.1.2 Non-Functional Requirements", "12", False),
        ("", "3.2 Feasibility Study", "13", False),
        ("", "3.2.1 Technical Feasibility", "13", False),
        ("", "3.2.2 Economic Feasibility", "14", False),
        ("", "3.2.3 Social and Ethical Feasibility", "14", False),
        ("", "3.3 System Specification", "15", False),
        ("", "3.3.1 Hardware Specification", "15", False),
        ("", "3.3.2 Software Specification", "15", False),
        ("4.", "DESIGN APPROACH AND DETAILS", "16", True),
        ("", "4.1 System Architecture", "16", False),
        ("", "4.2 Design", "16", False),
        ("", "4.2.1 Data Flow Diagram", "17", False),
        ("", "4.2.2 Use Case Diagram", "18", False),
        ("", "4.2.3 Class Diagram", "19", False),
        ("", "4.2.4 Sequence Diagram", "20", False),
        ("5.", "IMPLEMENTATION AND EXPERIMENTAL METHODOLOGY", "21", True),
        ("", "5.1 Implementation Environment and Specifications", "21", False),
        ("", "5.2 Five-Stage Attack Taxonomy and Operational Scope", "21", False),
        ("", "5.3 Dataset Construction and In-House Testbed", "22", False),
        ("", "5.4 The 31 Causal Feature Engineering Framework", "25", False),
        ("", "5.5 Machine Learning Baseline Models", "27", False),
        ("", "5.6 Hidden Markov Model (HMM) Transition Dynamics", "28", False),
        ("", "5.7 Causal Ensemble Architecture and Decision Logic", "28", False),
        ("", "5.8 Evaluation Metrics", "29", False),
        ("", "5.9 Experimental Setup and Reproducibility", "30", False),
        ("6.", "RESULTS AND EVALUATION", "31", True),
        ("", "6.1 Benchmark Comparison on Held-Out Test Set", "31", False),
        ("", "6.2 Per-Class Performance and Confusion Matrix Analysis", "32", False),
        ("", "6.3 Transition Edge Accuracy and Sequential Consistency", "33", False),
        ("", "6.4 False-Positive Rate Reduction via Temporal Context", "34", False),
        ("", "6.5 Session-Level Attack Campaign Detection", "34", False),
        ("", "6.6 Early Detection Latency & Defensive Lead Time", "35", False),
        ("", "6.7 Next-Stage Attack Intent Forecasting", "35", False),
        ("", "6.8 Real-Time Streaming and Inference Latency", "36", False),
        ("7.", "DISCUSSION", "38", True),
        ("", "7.1 The Value of Context in Web Application Security", "38", False),
        ("", "7.2 Error Analysis: The Fuzzing-Injection Boundary", "38", False),
        ("", "7.3 Operational Feasibility in Security Operations Centers", "38", False),
        ("", "7.4 Latency vs. Accuracy Trade-Off", "39", False),
        ("", "7.5 The Zero-Leakage Experimental Guarantee", "39", False),
        ("", "7.6 Threats to Validity and Limitations", "40", False),
        ("8.", "CONCLUSION AND FUTURE WORK", "41", True),
        ("", "8.1 Summary of Contributions", "41", False),
        ("", "8.2 Future Research Directions", "41", False),
        ("9.", "REFERENCES", "43", True),
    ]
    
    col_widths = [807085, 3780790, 1137285]
    
    for sl_no, title, page_no, is_major in toc_entries:
        row = toc.add_row()
        for idx, w in enumerate(col_widths):
            row.cells[idx].width = w
        format_cell(row.cells[0], sl_no, bold=is_major, size_pt=10.5, align=WD_ALIGN_PARAGRAPH.LEFT)
        format_cell(row.cells[1], title, bold=is_major, size_pt=10.5, align=WD_ALIGN_PARAGRAPH.LEFT)
        format_cell(row.cells[2], page_no, bold=is_major, size_pt=10.5, align=WD_ALIGN_PARAGRAPH.RIGHT)
        
    print(f"Table of Contents updated with {len(toc_entries)} entries.")

def enforce_pure_black_font(doc):
    print("Enforcing pure black font (#000000) and Times New Roman across entire document...")
    total_runs = 0
    # Paragraphs
    for p in doc.paragraphs:
        for r in p.runs:
            r.font.name = FONT_NAME
            r.font.color.rgb = BLACK
            total_runs += 1
            
    # Tables
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.name = FONT_NAME
                        r.font.color.rgb = BLACK
                        total_runs += 1
                        
    # Headers and Footers
    for s in doc.sections:
        for p in s.header.paragraphs:
            for r in p.runs:
                r.font.name = FONT_NAME
                r.font.color.rgb = BLACK
                total_runs += 1
        for p in s.footer.paragraphs:
            for r in p.runs:
                r.font.name = FONT_NAME
                r.font.color.rgb = BLACK
                total_runs += 1
                        
    print(f"Enforced black font on all {total_runs} text runs in the document.")

def main():
    if not os.path.exists(BACKUP_PATH):
        if os.path.exists(DOC_PATH):
            shutil.copyfile(DOC_PATH, BACKUP_PATH)
            print(f"Created backup at: {BACKUP_PATH}")
        else:
            raise FileNotFoundError(f"Neither {BACKUP_PATH} nor {DOC_PATH} exists!")
            
    print(f"Loading base document from backup: {BACKUP_PATH}")
    doc = docx.Document(BACKUP_PATH)
    print(f"Initial document loaded. Total paragraphs: {len(doc.paragraphs)}, Total tables: {len(doc.tables)}, Total body elements: {len(doc._body._element)}")
    
    # Step 1: Update Chapters 1 to 4
    apply_updates_ch1_to_4(doc)
    
    # Step 2: Dynamically locate '5. REFERENCES' and truncate everything from that index onwards
    trunc_idx = None
    for idx, el in enumerate(doc._body._element):
        tag = el.tag.split('}')[-1]
        if tag == 'p':
            p = docx.text.paragraph.Paragraph(el, doc)
            if '5. REFERENCES' in p.text or p.text.strip().startswith('5. REFERENCES'):
                trunc_idx = idx
                break
                
    if trunc_idx is None:
        raise ValueError("Could not find '5. REFERENCES' in document body elements!")
        
    print(f"Found '5. REFERENCES' at body element index {trunc_idx}. Truncating old trailing elements...")
    initial_elements = len(doc._body._element)
    while len(doc._body._element) > trunc_idx:
        doc._body._element.remove(doc._body._element[trunc_idx])
    print(f"Removed {initial_elements - len(doc._body._element)} old trailing elements.")
    
    # Step 3: Append Chapters 5, 6, 7, 8, and Chapter 9 References
    append_chapter_5(doc)
    append_chapter_6(doc)
    append_chapters_7_8_9(doc)
    
    # Step 4: Update Table of Contents (Table 2)
    update_table_of_contents(doc)
    
    # Step 5: Enforce pure black font across every run
    enforce_pure_black_font(doc)
    
    # Step 6: Save document
    print(f"Saving updated document to: {DOC_PATH}")
    doc.save(DOC_PATH)
    print("[SUCCESS] Document successfully generated and saved!")
    
    local_repo_doc = os.path.join(os.path.dirname(os.path.abspath(__file__)), "BCSE497J Project I Report.docx")
    if os.path.abspath(DOC_PATH) != os.path.abspath(local_repo_doc):
        shutil.copyfile(DOC_PATH, local_repo_doc)
        print(f"Copied updated document to local repo: {local_repo_doc}")
        
    # Step 7: Verify saved document
    doc_check = docx.Document(DOC_PATH)
    print(f"Verification: Document reloaded successfully!")
    print(f"Total Paragraphs: {len(doc_check.paragraphs)}")
    print(f"Total Tables: {len(doc_check.tables)}")
    print(f"Total Body Elements: {len(doc_check._body._element)}")

if __name__ == "__main__":
    main()
